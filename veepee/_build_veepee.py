"""Vente privée Veepee : construit le fichier unique (1 onglet) depuis _veepee_data.json (voir _extract_veepee.py)."""
import html
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
HERE = Path(__file__).parent
REMISE = 0.40
# EAN attribués le 13/10/2026 (absents d'Odoo à l'extraction)
EAN = {"C0191": "5413393004350", "C0192": "5413393004367", "C0193": "5413393004374"}
# RRP TTC qui corrigent le list_price Odoo (V0905 : 7,00 dans Odoo, 12,00 au tarif Noël)
RRP = {"V0905": 12.00}
# Lignes retirées de l'offre
EXCLUS = {"I070"}

# Famille + saveurs dominantes, FR | NL
SAVEURS = {
    "257": ("Infusion de fruits — pomme, citron, vanille", "Vruchteninfusie — appel, citroen, vanille"),
    "280": ("Thé vert — rose, poire", "Groene thee — roos, peer"),
    "340": ("Infusion épicée — réglisse, anis, fleur de mauve", "Kruideninfusie — zoethout, anijs, kaasjeskruid"),
    "483": ("Thé vert BIO — maté, citronnelle", "Groene thee BIO — maté, citroengras"),
    "601": ("Thé vert — goji, açaï, sureau", "Groene thee — goji, açaí, vlierbes"),
    "607": ("Thé vert — fraise, rose, canneberge", "Groene thee — aardbei, roos, veenbes"),
    "608": ("Thé vert — bergamote, bleuet", "Groene thee — bergamot, korenbloem"),
    "610": ("Thé noir — bergamote, citron, orange", "Zwarte thee — bergamot, citroen, sinaasappel"),
    "632": ("Rooibos — poire, datte, raisin, citronnelle", "Rooibos — peer, dadel, rozijn, citroengras"),
    "635": ("Maté — menthe poivrée, citronnelle, orange", "Maté — pepermunt, citroengras, sinaasappel"),
    "664": ("Maté vert nature", "Groene maté natuur"),
    "666": ("Thé noir nature — Inde", "Zwarte thee natuur — India"),
    "685": ("Thé noir nature BIO — Darjeeling", "Zwarte thee natuur BIO — Darjeeling"),
    "709": ("Infusion de fruits — ananas, pomme, citron, pêche", "Vruchteninfusie — ananas, appel, citroen, perzik"),
    "710": ("Infusion de fruits — mandarine, gingembre, fruit du dragon", "Vruchteninfusie — mandarijn, gember, drakenvrucht"),
    "728": ("Infusion de plantes — pomme, framboise, millepertuis", "Kruideninfusie — appel, framboos, sint-janskruid"),
    "737": ("Rooibos — caramel", "Rooibos — karamel"),
    "742": ("Oolong — fraise, goji", "Oolong — aardbei, goji"),
    "772": ("Thé noir — ananas, rose, bleuet", "Zwarte thee — ananas, roos, korenbloem"),
    "775": ("Thé noir — framboise, canneberge, amande, orange", "Zwarte thee — framboos, veenbes, amandel, sinaasappel"),
    "778": ("Thé noir — spéculoos : cardamome, amande, cannelle", "Zwarte thee — speculaas: kardemom, amandel, kaneel"),
    "788": ("Thé vert nature — Sencha du Japon", "Groene thee natuur — Japanse Sencha"),
    "793": ("Thé vert — orange, amande, fleur d'oranger", "Groene thee — sinaasappel, amandel, oranjebloesem"),
    "813": ("Thé vert — menthe", "Groene thee — munt"),
    "816": ("Thé vert BIO — orange, cannelle, cardamome", "Groene thee BIO — sinaasappel, kaneel, kardemom"),
    "821": ("Thé noir — ananas, mangue, papaye", "Zwarte thee — ananas, mango, papaja"),
    "822": ("Infusion de fruits — coco, melon", "Vruchteninfusie — kokos, meloen"),
    "826": ("Infusion de fruits BIO — coco, ananas", "Vruchteninfusie BIO — kokos, ananas"),
    "829": ("Thé vert & noir — ananas, orange, citron", "Groene & zwarte thee — ananas, sinaasappel, citroen"),
    "834": ("Maté — coco", "Maté — kokos"),
    "837": ("Thé vert — pistache, pomme, canneberge", "Groene thee — pistache, appel, veenbes"),
    "838": ("Oolong — framboise, rose", "Oolong — framboos, roos"),
    "843": ("Oolong BIO — lacté, vanillé", "Oolong BIO — melkachtig, vanille"),
    "847": ("Infusion de fruits — pastèque, menthe, sureau", "Vruchteninfusie — watermeloen, munt, vlierbes"),
    "848": ("Thé vert BIO — gingembre, citron", "Groene thee BIO — gember, citroen"),
    "866": ("Thé noir — cassis", "Zwarte thee — zwarte bes"),
    "870": ("Pu erh — églantier, maté, citron", "Pu-erh — rozenbottel, maté, citroen"),
    "875": ("Rooibos — orange, cassis, cannelle", "Rooibos — sinaasappel, zwarte bes, kaneel"),
    "879": ("Maté — mangue, clémentine, banane", "Maté — mango, clementine, banaan"),
    "885": ("Thé vert — coco, ananas", "Groene thee — kokos, ananas"),
    "889": ("Maté — grenade, citron", "Maté — granaatappel, citroen"),
    "890": ("Thé vert — ananas, mangue, menthe", "Groene thee — ananas, mango, munt"),
    "897": ("Pu erh BIO — orange sanguine, citronnelle, hibiscus", "Pu-erh BIO — bloedsinaasappel, citroengras, hibiscus"),
    "899": ("Infusion de fruits BIO — pomme, orange, fraise, mangue", "Vruchteninfusie BIO — appel, sinaasappel, aardbei, mango"),
    "900": ("Thé noir BIO — pêche, ananas, citronnelle", "Zwarte thee BIO — perzik, ananas, citroengras"),
    "905": ("Oolong, thé blanc & vert — vanille, amande, pomme", "Oolong, witte & groene thee — vanille, amandel, appel"),
    "C0191": ("Infusion de fruits — chocolat, noisette", "Vruchteninfusie — chocolade, hazelnoot"),
    "C0192": ("Rooibos — amande, ananas", "Rooibos — amandel, ananas"),
    "C0193": ("Infusion de fruits — canneberge, orange", "Vruchteninfusie — veenbes, sinaasappel"),
    "C0187": ("Coffret 25 infusettes : cacao-noisette, rooibos amande, orange-cannelle, oolong vanille, canneberge-orange",
              "Box 25 theezakjes: cacao-hazelnoot, rooibos amandel, sinaasappel-kaneel, oolong vanille, veenbes-sinaasappel"),
    "C0188": ("Coffret 4 × 50 g vrac + pince à thé : cacao-noisette, orange-cannelle, rooibos amande, canneberge-orange",
              "Box 4 × 50 g losse thee + theeklem: cacao-hazelnoot, sinaasappel-kaneel, rooibos amandel, veenbes-sinaasappel"),
    "C0185": ("Coffret 10 infusettes sans théine : vanille-cannelle, cerise-amaretto, menthe-rose-lavande, masala chai, canneberge-orange",
              "Doosje 10 theezakjes cafeïnevrij: vanille-kaneel, kers-amaretto, munt-roos-lavendel, masala chai, veenbes-sinaasappel"),
    "C0186": ("Coffret 10 infusettes de thé : spéculoos, cerise-jasmin, oolong vanille, orange-amande, pomme-cannelle",
              "Doosje 10 theezakjes thee: speculaas, kers-jasmijn, oolong vanille, sinaasappel-amandel, appel-kaneel"),
}


def saveurs(ref):
    return SAVEURS.get(ref) or SAVEURS.get(ref[2:].lstrip("0")) or SAVEURS.get(ref[1:].lstrip("0")) or ("", "")


def texte(h):
    """HTML Shopify → texte brut lisible (sans le JSON-LD), paragraphes conservés."""
    h = re.sub(r"<script.*?</script>|<!--.*?-->", "", h or "", flags=re.S)
    h = re.sub(r"</(p|h\d|li|div|ul)>|<br\s*/?>", "\n", h)
    h = re.sub(r"<li[^>]*>", "• ", h)
    t = html.unescape(re.sub(r"<[^>]+>", "", h))
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    return "\n".join(l.strip() for l in t.splitlines() if l.strip())


# Titres Shopify qui ne suivent pas le schéma « Famille saveurs - Nom - Teatower »
NOMS = {
    "V0905": ("Un Noël Vanillé", "Een Gevanielleerde Kerst"),
    "C0191": ("Trésor des lutins", "Schat van de elfen"),
    "C0192": ("Gourmandise de Noël", "Kerstlekkernij"),
    "C0193": ("Féérie d'Hiver", "Wintersprookje"),
    "C0187": ("Coffret découverte de Noël", "Kerst ontdekkingsbox"),
    "C0188": ("La hotte de Noël", "De kerstsok"),
    "C0185": ("La petite boîte à infusions", "Het kleine infusiedoosje"),
    "C0186": ("La petite boîte à thés", "Het kleine theedoosje"),
}


def nom(title, odoo_name):
    parts = [p.strip() for p in (title or "").split(" - ")]
    parts = [p for p in parts if p and not p.startswith("Teatower")]
    return parts[1] if len(parts) >= 2 else (parts[0] if parts else odoo_name)


def fmt(e):
    m = re.search(r"\(([^)]*(?:infusettes|vrac|g)[^)]*)\)", e.get("odoo_name") or "")
    if m:
        return m.group(1).strip()
    return {"C0191": "Boîte 50 g", "C0192": "Boîte 50 g", "C0193": "Boîte 50 g", "V0905": "100 g vrac",
            "C0187": "Coffret 25 infusettes", "C0188": "Coffret 4 × 50 g + pince", "C0185": "Boîte 10 infusettes",
            "C0186": "Boîte 10 infusettes"}.get(e["ref"], e.get("variant", ""))


def dlu_min(dlus):
    if not dlus:
        return ""
    return min(dlus, key=lambda d: datetime.strptime(d, "%d/%m/%Y"))


data = [e for e in json.loads((HERE / "_veepee_data.json").read_text(encoding="utf-8")) if e["ref"] not in EXCLUS]
wb = Workbook()
ws = wb.active
ws.title = "Stock Veepee"
COLS = [("Réf.", 9), ("EAN", 15), ("Nom FR", 26), ("Nom NL", 26), ("Format", 18), ("PCB", 7), ("Saveurs FR", 40), ("Saveurs NL", 40),
        ("Qté dispo", 9), ("DLU la plus proche", 12), ("Prix de vente conseillé TTC (RRP)", 13), ("TVA", 6), ("RRP HTVA", 10),
        ("Prix Veepee TTC (−40 %)", 12), ("Prix Veepee HTVA", 11),
        ("URL photo", 45), ("Fiche web", 40), ("Description FR", 60), ("Description NL", 60)]
ws.append([c for c, _ in COLS])
for i, (_, w) in enumerate(COLS, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
    c = ws.cell(1, i)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F3A5F")
    c.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[1].height = 32
ws.freeze_panes = "C2"

ORDRE = {"I": 0, "V": 1, "C": 2}
data.sort(key=lambda e: (ORDRE.get(e["ref"][0], 9), e["ref"]))

for e in data:
    sfr, snl = saveurs(e["ref"])
    ttc, tva = RRP.get(e["ref"], e.get("ttc")), e.get("tva")
    vp_ttc = round(ttc * (1 - REMISE), 2) if ttc else None
    vp_ht = round(vp_ttc / (1 + tva / 100), 2) if ttc else None
    nfr, nnl = NOMS.get(e["ref"]) or (nom(e.get("title_fr"), e.get("odoo_name", "")), nom(e.get("title_nl"), ""))
    ws.append([e["ref"], EAN.get(e["ref"]) or e.get("ean", ""), nfr, nnl,
               fmt(e), None, sfr, snl, e["qty"], dlu_min(e.get("dlus")), ttc, (tva / 100) if tva is not None else None,
               round(ttc / (1 + tva / 100), 2) if ttc else None, vp_ttc, vp_ht, e.get("photo", ""), e.get("url", ""),
               texte(e.get("body_fr")), texte(e.get("body_nl"))])
    r = ws.max_row
    for col in (11, 13, 14, 15):
        ws.cell(r, col).number_format = '#,##0.00 "€"'
    ws.cell(r, 12).number_format = "0%"
    for col in range(1, len(COLS) + 1):
        ws.cell(r, col).alignment = Alignment(vertical="top", wrap_text=col in (7, 8))

n = ws.max_row
ws.append([])
ws.append(["TOTAL", "", f"{n - 1} références", "", "", "", "", "", f"=SUM(I2:I{n})"])
ws.cell(ws.max_row, 1).font = Font(bold=True)
ws.cell(ws.max_row, 9).font = Font(bold=True)
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{n}"

out = HERE / "Veepee - stock disponible vente privee.xlsx"
wb.save(out)
print(out, n - 1, "références,", sum(e["qty"] for e in data), "unités")
